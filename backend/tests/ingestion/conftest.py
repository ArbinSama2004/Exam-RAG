"""Builders for real PDF and DOCX files, so the loaders are tested end to end.

Generating the files here keeps binary fixtures out of the repository and makes
the structure each test relies on visible in the test itself.
"""

import io
from collections.abc import Sequence

import docx
import pymupdf
import pytest

BODY_SIZE = 11.0
HEADING_SIZE = 22.0
SUBHEADING_SIZE = 15.0


@pytest.fixture
def make_pdf():
    """Build a PDF from (text, font size) lines, one list per page."""

    def _make(pages: Sequence[Sequence[tuple[str, float]]]) -> bytes:
        document = pymupdf.open()
        for lines in pages:
            page = document.new_page()
            y = 72.0
            for text, size in lines:
                page.insert_text((72, y), text, fontsize=size)
                y += size * 2

        data: bytes = document.tobytes()
        document.close()
        return data

    return _make


@pytest.fixture
def make_docx():
    """Build a DOCX from (text, style) paragraphs."""

    def _make(paragraphs: Sequence[tuple[str, str]], table: Sequence[Sequence[str]] | None = None):
        document = docx.Document()
        for text, style in paragraphs:
            document.add_paragraph(text, style=style or None)
        if table:
            docx_table = document.add_table(rows=len(table), cols=len(table[0]))
            for row_index, row in enumerate(table):
                for column_index, value in enumerate(row):
                    docx_table.cell(row_index, column_index).text = value
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    return _make


@pytest.fixture
def make_pdf_bold():
    """Build a PDF whose first line is bold at the same size as the body text."""

    def _make(bold_line: str, body_lines: Sequence[str]) -> bytes:
        document = pymupdf.open()
        page = document.new_page()
        page.insert_text((72, 72), bold_line, fontsize=BODY_SIZE, fontname="hebo")
        y = 72.0 + BODY_SIZE * 2
        for text in body_lines:
            page.insert_text((72, y), text, fontsize=BODY_SIZE, fontname="helv")
            y += BODY_SIZE * 2

        data: bytes = document.tobytes()
        document.close()
        return data

    return _make
